"""API CMCI F1 — fichas vulnerabilidad + socioeconómica (§3.1 §3.2 §10.1 + Fase 1).

Blueprint cmci_bp (url_prefix='/cmci') -> se registra UNA vez en api/__init__.py
sobre api_bp => rutas finales /api/cmci/*. NO usar prefijo /api aquí
(evita doble /api/api).

Persistencia primaria: modelos cmci.py (VulnerabilityAssessment,
SocioeconomicAssessment, ScoringParams vía ParamStore). POST crea siempre
una fila histórica (nunca overwrite); GET filtra en SQL (center/desde/hasta)
con paginación; priorización ordena en SQL (prioridad, alerta, total)
replicando Y/Z; dashboard agrega con COUNT/AVG SQL. Los dicts _VULN_DB/_SOC_DB
quedan SOLO como fallback si la DB no está disponible (try/except).

Incluye GET /vulnerability/<id>/ml-explain (F2 shadow: informa pero nunca
modifica el total determinista).
"""
import logging
from datetime import date, datetime

from flask import Blueprint, request, jsonify, g
from sqlalchemy import func

from middleware.tenant_context import tenant_required, TenantContext
from utils.role_helpers import (can_create_ficha, is_catering,
                                can_view_cmci_global, is_coordinadora,
                                is_admin_override, is_central_global)
from utils.phone_utils import normalize_phone_e164
from services.scoring import vulnerability_engine as VE
from services.scoring import socioeconomic_engine as SE
from services.scoring.params_store import ParamStore, load_seed
from services.scoring.explain import explain_vulnerability, explain_socioeconomic

logger = logging.getLogger(__name__)

cmci_bp = Blueprint("cmci", __name__, url_prefix="/cmci")

STORE = ParamStore()  # seed v1_validada_2026-09-25 + overrides por scope (DB-first)

# Fallback histórico en memoria (un registro por valoración, nunca overwrite).
# SOLO se usa si la DB no está disponible; la vía principal es SQL.
_VULN_DB = {}
_SOC_DB = {}
_SEQ = {"vuln": 0, "soc": 0}

# Clave reservada dentro de los JSON para metadatos del contrato API
# (center/child_name/observations/params_version/source + payload íntegro)
# que no tienen columna nativa en los modelos.
_META = "__meta__"

ADMIN_ROLES = ("license_admin", "super_admin")  # PUT /params solo license_admin


def _tenant():
    return TenantContext.get_current_tenant_id()


def _now_iso():
    return datetime.now().isoformat(timespec="seconds")


def _rollback_quiet():
    try:
        from models import db
        db.session.rollback()
    except Exception:
        pass


def _require_admin():
    user = getattr(g, "current_user", None)
    role = getattr(getattr(user, "role", None), "name", None)
    if role not in ADMIN_ROLES:
        return jsonify({"error": "Solo license_admin puede modificar parámetros"}), 403
    return None


def _scope():
    return request.args.get("scope", "global")


def _check_idor(record):
    """IDOR: el registro debe pertenecer al tenant (cmci==tenant)."""
    if not record or record.get("tenant_id") != _tenant():
        return jsonify({"error": "Registro no encontrado"}), 404
    return None


def _deny_unless_can_create():
    """F5 RBAC §10.7 (sin firma electrónica): Educadora crea, Auxiliar
    apoya; Coordinadora revisa en papel impreso; Central solo lectura;
    Catering solo menú/ingesta (403 aquí)."""
    user = getattr(g, "current_user", None)
    if not can_create_ficha(user):
        if is_catering(user):
            return jsonify({"error": "Rol catering: solo menú e ingesta diaria"}), 403
        return jsonify({"error": "Su rol revisa en papel impreso; "
                                 "la creación es de Educadora/Auxiliar"}), 403
    return None


def _deny_unless_can_edit():
    """Edición de fichas: Educadora/Auxiliar (crean) + Coordinadora
    (revisa y corrige) + admin override. Central solo lectura,
    Catering solo menú/ingesta."""
    user = getattr(g, "current_user", None)
    if can_create_ficha(user) or is_coordinadora(user) or is_admin_override(user):
        return None
    if is_catering(user):
        return jsonify({"error": "Rol catering: solo menú e ingesta diaria"}), 403
    return jsonify({"error": "Su rol es solo lectura en fichas"}), 403


def _deny_unless_can_upload_biblioteca():
    """Biblioteca: solo propietario (license_admin/super_admin) o
    coordinador principal (coordinator) pueden subir plantillas.
    Resto de roles: solo descarga (ver GET/download sin restricción)."""
    user = getattr(g, "current_user", None)
    if is_admin_override(user) or is_coordinadora(user):
        return None
    # license_admin/super_admin ya cubiertos por is_admin_override;
    # coordinador general / central con scope global también puede subir
    if is_central_global(user):
        return None
    return jsonify({"error": "Solo el propietario o el coordinador principal "
                             "pueden cargar documentos en la biblioteca"}), 403


def _sync_child_family_from_ficha(body, answers_or_data):
    """Propaga identificación de la ficha al Child/Family/Representative.

    body: {child_id, ...}; answers_or_data: dict con claves opcionales
    {nacimiento/birth_date, sexo/gender, representante, parentesco,
     telefono/phone, direccion/address, sector, cedula}.
    Retorna (child_id | None). Nunca rompe la creación de la ficha:
    errores de sync se loguean y se ignoran.
    """
    try:
        from models import db
        from models.child import Child, Family, Representative

        child_id = body.get("child_id") if isinstance(body, dict) else None
        if not child_id:
            return None
        try:
            child_id = int(child_id)
        except (TypeError, ValueError):
            return None
        tenant_id = _tenant()
        child = Child.query.filter_by(id=child_id, tenant_id=tenant_id).first()
        if not child:
            # multi-centro global: permite vincular dentro de la licencia
            user = getattr(g, "current_user", None)
            if can_view_cmci_global(user):
                child = Child.query.filter_by(id=child_id).first()
            if not child:
                return None
        payload = answers_or_data if isinstance(answers_or_data, dict) else {}
        ident = payload.get("identificacion") if isinstance(payload.get("identificacion"), dict) else payload

        def _pick(*keys):
            for k in keys:
                v = ident.get(k)
                if v not in (None, ""):
                    return v
            return None

        nacimiento = _pick("nacimiento", "birth_date", "fecha_nacimiento")
        if nacimiento:
            day, _ = _parse_day(nacimiento)
            if day is not None:
                child.birth_date = day
        sexo = _pick("sexo", "gender")
        if sexo:
            s = str(sexo).strip().lower()
            if s.startswith("m"):
                child.gender = "masculino"
            elif s.startswith("f"):
                child.gender = "femenino"
        ced = _pick("cedula", "cedula_nino")
        if ced:
            child.cedula = str(ced).strip()[:10]
        db.session.add(child)

        family = getattr(child, "family", None)
        if family is None and getattr(child, "family_id", None):
            family = Family.query.filter_by(id=child.family_id).first()
        if family is not None:
            addr = _pick("direccion", "address", "domicilio")
            if addr:
                family.address = str(addr)
            sec = _pick("sector")
            if sec:
                family.sector = str(sec)
            tel = _pick("telefono", "teléfono", "phone", "phone_primary", "celular")
            if tel:
                family.phone_primary = str(tel).strip()[:20]
            # socioeconómica -> Family
            per_cap = payload.get("per_capita") if "per_capita" in payload else _pick("per_capita", "perCapita")
            if per_cap is None and isinstance(payload.get("resultado"), dict):
                per_cap = payload["resultado"].get("per_capita")
            try:
                if per_cap is not None and float(per_cap) >= 0:
                    family.vulnerability_score = int(float(per_cap))
            except (TypeError, ValueError):
                pass
            db.session.add(family)

            rep_name = _pick("representante", "representante_nombre")
            parentesco = _pick("parentesco", "relationship")
            rep_tel = _pick("telefono", "teléfono", "phone", "celular")
            if rep_name:
                parts = str(rep_name).strip().split()
                first = parts[0] if parts else str(rep_name)
                last = " ".join(parts[1:]) if len(parts) > 1 else "-"
                existing = Representative.query.filter_by(family_id=family.id).first()
                rel = str(parentesco or "otro").strip().lower()
                valid_rels = {"madre", "padre", "abuelo", "abuela", "tio", "tia",
                              "tutor_legal", "otro"}
                if rel not in valid_rels:
                    rel = "madre" if "madr" in rel else ("padre" if "padr" in rel else "otro")
                if existing:
                    existing.first_name = first[:100]
                    existing.last_name = last[:100]
                    existing.relationship = rel
                    if rep_tel:
                        existing.phone = str(rep_tel).strip()[:20]
                    db.session.add(existing)
        db.session.commit()
        return child.id
    except Exception as exc:
        logger.warning("sync_child_family: no se pudo propagar (%s)", exc)
        _rollback_quiet()
        return None


def _enforce_site(requested_center):
    """F5 filtro site forzado: si el usuario no es global y pide otro
    centro distinto al suyo → 404 (IDOR tenant==site)."""
    if not requested_center:
        return None
    user = getattr(g, "current_user", None)
    if can_view_cmci_global(user):
        return None
    own = TenantContext.get_current_cmci_code()
    if own and str(requested_center) != str(own):
        return jsonify({"error": "Registro no encontrado"}), 404
    return None


_PHONE_KEYS = ("telefono", "teléfono", "phone", "whatsapp", "celular")


def _normalize_phones(payload):
    """F5: depuración 09→593 (E.164 por país). No rompe si no hay teléfono."""
    if not isinstance(payload, dict):
        return payload
    iso = TenantContext.get_current_country_iso()
    for key in _PHONE_KEYS:
        if payload.get(key):
            normalized = normalize_phone_e164(payload[key], iso)
            if normalized:
                payload[key] = normalized
    return payload


def _paginate(items, page, per_page):
    total = len(items)
    start = (page - 1) * per_page
    return {"data": items[start:start + per_page], "page": page,
            "per_page": per_page, "total": total}


def _page_args(args):
    page = max(int(args.get("page", 1)), 1)
    per_page = min(max(int(args.get("per_page", 20)), 1), 100)
    return page, per_page


def _in_range_date(value, _from, _to):
    if not value:
        return False
    day = str(value)[:10]
    if _from and day < _from:
        return False
    if _to and day > _to:
        return False
    return True


def _parse_day(value):
    """Normaliza fecha ISO/date/datetime → (date|None, raw_iso|None)."""
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime):
        return value.date(), value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value, value.isoformat()
    text = str(value).strip()
    if not text:
        return None, None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date(), text
    except ValueError:
        pass
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date(), text
    except ValueError:
        return None, text


# --------------------------- VULNERABILIDAD (DB) ---------------------------

def _vuln_insert_db(body, answers, result, params_version, source):
    """Inserta fila histórica en vulnerability_assessments (nunca overwrite)."""
    from models import db
    from models.cmci import VulnerabilityAssessment

    body = body or {}
    answers = dict(answers) if isinstance(answers, dict) else {}
    result = result or {}
    assessed_in = body.get("assessed_at") or answers.get("assessed_at")
    day, raw = _parse_day(assessed_in or _now_iso())
    birth_in = answers.get("birth_date")
    bday, _ = _parse_day(birth_in) if birth_in else (None, None)
    stored_answers = dict(answers)
    stored_answers[_META] = {
        "center": body.get("center"), "child_name": body.get("child_name"),
        "observations": body.get("observations"),
        "params_version": params_version, "source": source,
        "assessed_at": raw, "result": result,
    }
    row = VulnerabilityAssessment(
        child_id=body.get("child_id"), center_id=_tenant(),
        code=body.get("code"), assessed_at=day, birth_date=bday,
        age_months=result.get("age_months"),
        in_range=(result.get("age_range") != "FUERA DE RANGO"),
        answers=stored_answers, scores=result.get("scores") or {},
        subtotals=result.get("subtotals") or {},
        total=result.get("total") or 0.0, level=result.get("level"),
        semaphore=result.get("semaphore"),
        protection_alert=bool(result.get("protection_flag")),
        priority=result.get("priority"),
        childcare_need=result.get("childcare_need"),
        alerts=result.get("alerts") or {},
        status=body.get("status", "En proceso"),
        created_by=getattr(getattr(g, "current_user", None), "id", None),
    )
    db.session.add(row)
    db.session.commit()
    return _vuln_to_record(row)


def _vuln_to_record(row):
    """Serializa un row al contrato API (idéntico al histórico en memoria)."""
    answers = dict(getattr(row, "answers", None) or {})
    meta = answers.pop(_META, {})
    if not isinstance(meta, dict):
        meta = {}
    result = meta.get("result")
    if not isinstance(result, dict):
        total = getattr(row, "total", 0.0) or 0.0
        prot = bool(getattr(row, "protection_alert", False))
        result = {
            "total": total, "pct": total / 100.0,
            "level": getattr(row, "level", None),
            "priority": getattr(row, "priority", None),
            "protection_flag": prot,
            "protection_alert": "SÍ" if prot else "NO",
            "semaphore": getattr(row, "semaphore", None),
            "childcare_need": getattr(row, "childcare_need", None),
            "scores": getattr(row, "scores", None) or {},
            "subtotals": getattr(row, "subtotals", None) or {},
            "contrib": {}, "alerts": getattr(row, "alerts", None) or {},
            "age_months": getattr(row, "age_months", None),
            "order_y": VE.order_y(getattr(row, "priority", None) or "PRIORIDAD 3",
                                  prot, total),
        }
    assessed = meta.get("assessed_at")
    if not assessed and getattr(row, "assessed_at", None) is not None:
        assessed = row.assessed_at.isoformat()
    return {
        "id": row.id, "tenant_id": getattr(row, "center_id", None),
        "child_id": getattr(row, "child_id", None),
        "center": meta.get("center"), "code": getattr(row, "code", None),
        "child_name": meta.get("child_name"), "status": getattr(row, "status", None),
        "assessed_at": assessed, "observations": meta.get("observations"),
        "answers": answers, "result": result,
        "params_version": meta.get("params_version"),
        "source": meta.get("source"),
        "created_by": getattr(row, "created_by", None),
    }


def _vuln_base_query(args):
    """Query base con filtros tenant + center/desde/hasta (+child_id/status)."""
    from models.cmci import VulnerabilityAssessment

    q = VulnerabilityAssessment.query.filter(
        VulnerabilityAssessment.center_id == _tenant())
    center = args.get("center")
    if center:
        from models.tenant import Tenant
        q = (q.join(Tenant,
                    VulnerabilityAssessment.center_id == Tenant.id)
              .filter(Tenant.cmci_code == str(center)))
    if args.get("child_id") is not None and str(args.get("child_id")) != "":
        try:
            q = q.filter(VulnerabilityAssessment.child_id == int(args["child_id"]))
        except (TypeError, ValueError):
            from models import db
            q = q.filter(db.false())  # id no numérico → sin resultados
    if args.get("status"):
        q = q.filter(VulnerabilityAssessment.status == args["status"])
    if args.get("from"):
        day, _ = _parse_day(args["from"])
        if day is not None:
            q = q.filter(VulnerabilityAssessment.assessed_at >= day)
    if args.get("to"):
        day, _ = _parse_day(args["to"])
        if day is not None:
            q = q.filter(VulnerabilityAssessment.assessed_at <= day)
    return q


def _vuln_list_db(args):
    from models.cmci import VulnerabilityAssessment

    page, per_page = _page_args(args)
    q = _vuln_base_query(args)
    total = q.count()
    rows = (q.order_by(VulnerabilityAssessment.assessed_at.desc(),
                       VulnerabilityAssessment.id.desc())
             .limit(per_page).offset((page - 1) * per_page).all())
    return {"data": [_vuln_to_record(r) for r in rows], "page": page,
            "per_page": per_page, "total": total}


def _vuln_detail_db(rec_id):
    from models.cmci import VulnerabilityAssessment

    row = (VulnerabilityAssessment.query
           .filter(VulnerabilityAssessment.id == rec_id,
                   VulnerabilityAssessment.center_id == _tenant()).first())
    return _vuln_to_record(row) if row is not None else None


def _vuln_rank_db(args):
    """Priorización: orden Y/Z en SQL (prioridad, alerta, total)."""
    from models.cmci import VulnerabilityAssessment

    page, per_page = _page_args(args)
    # Replica BASE col Y = priNum*100000 + alerta*1000 + total:
    # ordenar por Y desc equivale a (prioridad, alerta, total) desc, ya que
    # alerta*1000+total < 100000. Los literales 'PRIORIDAD 1/2/3' ordenan
    # ascendentemente igual que su rango, por eso prioridad ASC.
    q = (_vuln_base_query(args)
         .order_by(func.coalesce(VulnerabilityAssessment.priority,
                                 "PRIORIDAD 9").asc(),
                   VulnerabilityAssessment.protection_alert.desc(),
                   VulnerabilityAssessment.total.desc(),
                   VulnerabilityAssessment.id.asc()))
    total = q.count()
    rows = q.limit(per_page).offset((page - 1) * per_page).all()
    data = []
    for pos, row in enumerate(rows, start=(page - 1) * per_page + 1):
        rec = _vuln_to_record(row)
        res = rec["result"]
        data.append({"id": rec["id"], "code": rec.get("code"),
                     "child_id": rec.get("child_id"),
                     "child_name": rec.get("child_name"),
                     "center": rec.get("center"),
                     "assessed_at": rec.get("assessed_at"),
                     "age_months": res.get("age_months"),
                     "total": res.get("total"), "level": res.get("level"),
                     "priority": res.get("priority"),
                     "protection_alert": res.get("protection_alert"),
                     "semaphore": res.get("semaphore"),
                     "status": rec.get("status"),
                     "order_y": res.get("order_y"), "rank_z": pos})
    return {"data": data, "page": page, "per_page": per_page, "total": total}


def _vuln_dashboard_db(args):
    """Dashboard con agregados SQL (COUNT/AVG/GROUP BY)."""
    from models.cmci import VulnerabilityAssessment

    q = _vuln_base_query(args)
    id_col = VulnerabilityAssessment.id
    total = q.with_entities(func.count(id_col)).scalar() or 0
    avg = q.with_entities(func.avg(VulnerabilityAssessment.total)).scalar()
    by_level = {lvl: {"count": cnt, "pct": cnt / total if total else 0}
                for lvl, cnt in
                q.with_entities(VulnerabilityAssessment.level,
                                func.count(id_col))
                 .group_by(VulnerabilityAssessment.level).all()}
    by_prio = {prio: cnt for prio, cnt in
               q.with_entities(VulnerabilityAssessment.priority,
                               func.count(id_col))
                .group_by(VulnerabilityAssessment.priority).all()}
    alerts = (q.filter(VulnerabilityAssessment.protection_alert.is_(True))
               .with_entities(func.count(id_col)).scalar() or 0)
    if args.get("center"):
        by_center = {args["center"]: total}
    else:
        from models.tenant import Tenant
        center_col = func.coalesce(Tenant.cmci_code, "Sin centro")
        by_center = {c: cnt for c, cnt in
                     q.outerjoin(Tenant, VulnerabilityAssessment.center_id
                                 == Tenant.id)
                      .with_entities(center_col, func.count(id_col))
                      .group_by(center_col).all()}
    return {"total_E5": total, "avg_I5": (avg / 100.0 if avg is not None else None),
            "by_level": by_level, "by_priority_D17_D19": by_prio,
            "protection_alerts_E22": alerts, "by_center": by_center}


def _vuln_create_mem(body, answers, result):
    _SEQ["vuln"] += 1
    rec_id = _SEQ["vuln"]
    record = {
        "id": rec_id,
        "tenant_id": _tenant(),
        "child_id": body.get("child_id"),
        "center": body.get("center"),
        "code": body.get("code"),
        "child_name": body.get("child_name"),
        "status": body.get("status", "En proceso"),
        "assessed_at": body.get("assessed_at") or answers.get("assessed_at") or _now_iso(),
        "observations": body.get("observations"),
        "answers": answers,
        "result": result,
        "params_version": STORE.version,
        "source": "api",
        "created_by": getattr(getattr(g, "current_user", None), "id", None),
    }
    _VULN_DB[rec_id] = record
    return record


# --------------------------- VULNERABILIDAD ---------------------------

@cmci_bp.route("/vulnerability/compute", methods=["POST"])
@tenant_required
def vuln_compute():
    """Preview sin guardar: replica VALORACIÓN F78-F84 + G88-G97."""
    body = request.get_json(force=True, silent=True) or {}
    answers = body.get("answers", body)
    try:
        result = VE.compute(answers, STORE.get(body.get("scope", "global")))
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    result["explain"] = explain_vulnerability(result, answers, STORE.get(body.get("scope", "global")))
    return jsonify(result), 200


@cmci_bp.route("/vulnerability", methods=["POST"])
@tenant_required
def vuln_create():
    """Guarda NUEVO registro histórico (macro Guardar; nunca sobreescribe)."""
    denied = _deny_unless_can_create()
    if denied:
        return denied
    body = request.get_json(force=True, silent=True) or {}
    site_err = _enforce_site(body.get("center"))
    if site_err:
        return site_err
    answers = _normalize_phones(body.get("answers", {}))
    try:
        params = STORE.get(body.get("scope", "global"))
        result = VE.compute(answers, params)
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    try:
        record = _vuln_insert_db(body, answers, result, STORE.version, "api")
    except Exception as exc:
        logger.warning("vuln_create: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
        record = _vuln_create_mem(body, answers, result)
    # Vinculación niño: propaga identificación a Child/Family/Representative
    # y recalcula puntaje ya hecho arriba (result). No rompe si falla.
    _sync_child_family_from_ficha(body, answers)
    return jsonify(record), 201


@cmci_bp.route("/vulnerability", methods=["GET"])
@tenant_required
def vuln_list():
    """GET /vulnerability?child_id&center&status&from&to&page&per_page."""
    args = request.args
    site_err = _enforce_site(args.get("center"))
    if site_err:
        return site_err
    try:
        return jsonify(_vuln_list_db(args)), 200
    except Exception as exc:
        logger.warning("vuln_list: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
    items = [r for r in _VULN_DB.values() if r.get("tenant_id") == _tenant()]
    if args.get("child_id"):
        items = [r for r in items if str(r.get("child_id")) == args["child_id"]]
    if args.get("center"):
        items = [r for r in items if r.get("center") == args["center"]]
    if args.get("status"):
        items = [r for r in items if r.get("status") == args["status"]]
    if args.get("from") or args.get("to"):
        items = [r for r in items if _in_range_date(r.get("assessed_at"), args.get("from"), args.get("to"))]
    items.sort(key=lambda r: r.get("assessed_at", ""), reverse=True)
    page = max(int(args.get("page", 1)), 1)
    per_page = min(max(int(args.get("per_page", 20)), 1), 100)
    return jsonify(_paginate(items, page, per_page)), 200


@cmci_bp.route("/vulnerability/<int:rec_id>", methods=["GET"])
@tenant_required
def vuln_detail(rec_id):
    try:
        record = _vuln_detail_db(rec_id)
    except Exception as exc:
        logger.warning("vuln_detail: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
        err = _check_idor(_VULN_DB.get(rec_id))
        if err:
            return err
        return jsonify(_VULN_DB[rec_id]), 200
    if record is None:
        return jsonify({"error": "Registro no encontrado"}), 404
    return jsonify(record), 200


@cmci_bp.route("/vulnerability/<int:assessment_id>/ml-explain", methods=["GET"])
@tenant_required
def ml_explain(assessment_id: int):
    """Explica el shadow ML de una ficha de vulnerabilidad CMCI.

    Retorna {ml_priority_proba, top3, determinista_total, delta,
    ml_version, mode}. Si no hay modelo o la ficha no tiene proba
    persistida, recalcula shadow al vuelo (sin persistir ni alterar nada).
    El ML es shadow: NUNCA modifica el total determinista (delta 0.0).
    """
    try:
        from models.cmci import VulnerabilityAssessment
        from services.ml_calibrator import calibrate, explain_for_assessment

        rec = VulnerabilityAssessment.query.get(assessment_id)
        if not rec:
            return jsonify({'success': False,
                            'error': 'Ficha no encontrada'}), 404
        # F5 IDOR: la ficha debe pertenecer al tenant (site==tenant).
        tenant_id = TenantContext.get_current_tenant_id()
        if getattr(rec, 'center_id', None) not in (None, tenant_id):
            return jsonify({'success': False,
                            'error': 'Ficha no encontrada'}), 404

        proba = getattr(rec, 'ml_proba', None)
        ml_version = getattr(rec, 'ml_version', None)
        top3 = None
        if proba is None:
            # Recálculo shadow al vuelo (solo lectura, sin persistir)
            fd = dict(getattr(rec, 'answers', None) or {})
            fd.pop(_META, None)
            fd['subtotals'] = getattr(rec, 'subtotals', None) or {}
            fd['scores'] = getattr(rec, 'scores', None) or {}
            fd['alerts'] = getattr(rec, 'alerts', None) or {}
            _, info = calibrate(None, fd, getattr(rec, 'total', 0.0) or 0.0)
            proba = info.get('proba')
            top3 = info.get('top3', [])
            ml_version = info.get('ml_version')
        payload = explain_for_assessment(getattr(rec, 'total', 0.0) or 0.0,
                                         proba, top3=top3, ml_version=ml_version)
        payload['success'] = True
        payload['assessment_id'] = assessment_id
        return jsonify(payload), 200
    except Exception as e:
        logger.error(f'ml-explain fallo: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@cmci_bp.route("/priorizacion", methods=["GET"])
@tenant_required
def priorizacion():
    """Espejo PRIORIZACIÓN: orden Y/Z + paginación + filtros center/from/to."""
    args = request.args
    site_err = _enforce_site(args.get("center"))
    if site_err:
        return site_err
    try:
        return jsonify(_vuln_rank_db(args)), 200
    except Exception as exc:
        logger.warning("priorizacion: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
    items = [r for r in _VULN_DB.values() if r.get("tenant_id") == _tenant()]
    if args.get("center"):
        items = [r for r in items if r.get("center") == args["center"]]
    if args.get("from") or args.get("to"):
        items = [r for r in items if _in_range_date(r.get("assessed_at"), args.get("from"), args.get("to"))]
    ranked = VE.ranking([{"id": r["id"], "order_y": r["result"]["order_y"]} for r in items])
    rank_by_id = {x["id"]: x["rank_z"] for x in ranked}
    rows = []
    for r in items:
        res = r["result"]
        rows.append({"id": r["id"], "code": r.get("code"), "child_id": r.get("child_id"),
                     "child_name": r.get("child_name"), "center": r.get("center"),
                     "assessed_at": r.get("assessed_at"), "age_months": res.get("age_months"),
                     "total": res["total"], "level": res["level"],
                     "priority": res["priority"], "protection_alert": res["protection_alert"],
                     "semaphore": res["semaphore"], "status": r.get("status"),
                     "order_y": res["order_y"], "rank_z": rank_by_id[r["id"]]})
    rows.sort(key=lambda x: x["rank_z"])
    page = max(int(args.get("page", 1)), 1)
    per_page = min(max(int(args.get("per_page", 20)), 1), 100)
    return jsonify(_paginate(rows, page, per_page)), 200


@cmci_bp.route("/dashboard", methods=["GET"])
@tenant_required
def dashboard():
    """Conteos DASHBOARD: E5/I5/D9-D13(+pct)/D17-D19/E22/D25-D28."""
    args = request.args
    site_err = _enforce_site(args.get("center"))
    if site_err:
        return site_err
    try:
        return jsonify(_vuln_dashboard_db(args)), 200
    except Exception as exc:
        logger.warning("dashboard: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
    items = [r for r in _VULN_DB.values() if r.get("tenant_id") == _tenant()]
    if args.get("center"):
        items = [r for r in items if r.get("center") == args["center"]]
    if args.get("from") or args.get("to"):
        items = [r for r in items if _in_range_date(r.get("assessed_at"), args.get("from"), args.get("to"))]
    total = len(items)  # E5 COUNTIF(BASE!A,'?*')
    avg = (sum(r["result"]["total"] for r in items) / total / 100.0) if total else None  # I5
    levels = {}
    for r in items:
        levels[r["result"]["level"]] = levels.get(r["result"]["level"], 0) + 1
    by_level = {k: {"count": v, "pct": v / total if total else 0}  # D9:D13 + E9:E13
                for k, v in levels.items()}
    prios = {}
    for r in items:
        prios[r["result"]["priority"]] = prios.get(r["result"]["priority"], 0) + 1  # D17:D19
    alerts = sum(1 for r in items if r["result"]["protection_flag"])  # E22
    by_center = {}
    for r in items:
        by_center[r.get("center") or "Sin centro"] = by_center.get(r.get("center") or "Sin centro", 0) + 1  # D25:D28
    return jsonify({"total_E5": total, "avg_I5": avg, "by_level": by_level,
                    "by_priority_D17_D19": prios, "protection_alerts_E22": alerts,
                    "by_center": by_center}), 200


# --------------------------- SOCIOECONÓMICA (DB) ---------------------------

def _soc_insert_db(body, data, result, params_version, source):
    """Inserta fila histórica en socioeconomic_assessments (nunca overwrite)."""
    from models import db
    from models.cmci import SocioeconomicAssessment

    body = body or {}
    data = dict(data) if isinstance(data, dict) else {}
    result = result or {}
    assessed_in = body.get("assessed_at") or data.get("assessed_at")
    day, raw = _parse_day(assessed_in or _now_iso())
    ingresos = data.get("ingresos")
    egresos = data.get("egresos")
    stored_in = dict(ingresos) if isinstance(ingresos, dict) else {}
    stored_in[_META] = {
        "center": body.get("center"), "child_name": body.get("child_name"),
        "code": body.get("code"), "params_version": params_version,
        "source": source, "assessed_at": raw,
        "data": data, "result": result,
    }
    row = SocioeconomicAssessment(
        child_id=body.get("child_id"), center_id=_tenant(),
        code=body.get("code"), assessed_at=day,
        incomes=stored_in,
        expenses=dict(egresos) if isinstance(egresos, dict) else {},
        per_capita=result.get("per_capita"),
        expense_ratio=result.get("ratio_gasto"),
        coverage_services=result.get("cobertura_servicios"),
        dependency=result.get("personas_por_perceptor"),
        subscores=result.get("subscores") or {},
        total=result.get("total") or 0.0,
        classification=result.get("classification"),
    )
    db.session.add(row)
    db.session.commit()
    return _soc_to_record(row)


def _soc_to_record(row):
    """Serializa un row al contrato API (idéntico al histórico en memoria)."""
    incomes = dict(getattr(row, "incomes", None) or {})
    meta = incomes.pop(_META, {})
    if not isinstance(meta, dict):
        meta = {}
    data = meta.get("data")
    if not isinstance(data, dict):
        data = {"ingresos": incomes,
                "egresos": dict(getattr(row, "expenses", None) or {})}
    result = meta.get("result")
    if not isinstance(result, dict):
        result = {"per_capita": getattr(row, "per_capita", None),
                  "total": getattr(row, "total", 0.0) or 0.0,
                  "classification": getattr(row, "classification", None),
                  "subscores": getattr(row, "subscores", None) or {}}
    assessed = meta.get("assessed_at")
    if not assessed and getattr(row, "assessed_at", None) is not None:
        assessed = row.assessed_at.isoformat()
    return {
        "id": row.id, "tenant_id": getattr(row, "center_id", None),
        "child_id": getattr(row, "child_id", None),
        "center": meta.get("center"),
        "code": getattr(row, "code", None) or meta.get("code"),
        "child_name": meta.get("child_name"), "assessed_at": assessed,
        "data": data, "result": result,
        "params_version": meta.get("params_version"),
        "source": meta.get("source"),
    }


def _soc_base_query(args):
    from models.cmci import SocioeconomicAssessment

    q = SocioeconomicAssessment.query.filter(
        SocioeconomicAssessment.center_id == _tenant())
    if args.get("center"):
        from models.tenant import Tenant
        q = (q.join(Tenant,
                    SocioeconomicAssessment.center_id == Tenant.id)
              .filter(Tenant.cmci_code == str(args["center"])))
    if args.get("child_id") is not None and str(args.get("child_id")) != "":
        try:
            q = q.filter(SocioeconomicAssessment.child_id == int(args["child_id"]))
        except (TypeError, ValueError):
            from models import db
            q = q.filter(db.false())
    if args.get("from"):
        day, _ = _parse_day(args["from"])
        if day is not None:
            q = q.filter(SocioeconomicAssessment.assessed_at >= day)
    if args.get("to"):
        day, _ = _parse_day(args["to"])
        if day is not None:
            q = q.filter(SocioeconomicAssessment.assessed_at <= day)
    return q


def _soc_list_db(args):
    from models.cmci import SocioeconomicAssessment

    page, per_page = _page_args(args)
    q = _soc_base_query(args)
    total = q.count()
    rows = (q.order_by(SocioeconomicAssessment.assessed_at.desc(),
                       SocioeconomicAssessment.id.desc())
             .limit(per_page).offset((page - 1) * per_page).all())
    return {"data": [_soc_to_record(r) for r in rows], "page": page,
            "per_page": per_page, "total": total}


def _soc_detail_db(rec_id):
    from models.cmci import SocioeconomicAssessment

    row = (SocioeconomicAssessment.query
           .filter(SocioeconomicAssessment.id == rec_id,
                   SocioeconomicAssessment.center_id == _tenant()).first())
    return _soc_to_record(row) if row is not None else None


def _soc_create_mem(body, data, result):
    _SEQ["soc"] += 1
    rec_id = _SEQ["soc"]
    record = {"id": rec_id, "tenant_id": _tenant(), "child_id": body.get("child_id"),
              "center": body.get("center"), "code": body.get("code"),
              "child_name": body.get("child_name"),
              "assessed_at": body.get("assessed_at") or data.get("assessed_at") or _now_iso(),
              "data": data, "result": result, "params_version": STORE.version,
              "source": "api",
              "created_by": getattr(getattr(g, "current_user", None), "id", None)}
    _SOC_DB[rec_id] = record
    return record


# --------------------------- SOCIOECONÓMICA (espejo) ---------------------------

@cmci_bp.route("/socioeconomic/compute", methods=["POST"])
@tenant_required
def socio_compute():
    body = request.get_json(force=True, silent=True) or {}
    data = body.get("data", body)
    try:
        result = SE.compute(data, STORE.get(body.get("scope", "global")))
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    result["explain"] = explain_socioeconomic(result)
    return jsonify(result), 200


@cmci_bp.route("/socioeconomic", methods=["POST"])
@tenant_required
def socio_create():
    denied = _deny_unless_can_create()
    if denied:
        return denied
    body = request.get_json(force=True, silent=True) or {}
    site_err = _enforce_site(body.get("center"))
    if site_err:
        return site_err
    data = _normalize_phones(body.get("data", {}))
    try:
        result = SE.compute(data, STORE.get(body.get("scope", "global")))
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    try:
        record = _soc_insert_db(body, data, result, STORE.version, "api")
    except Exception as exc:
        logger.warning("socio_create: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
        record = _soc_create_mem(body, data, result)
    _sync_child_family_from_ficha(body, data)
    return jsonify(record), 201


@cmci_bp.route("/socioeconomic", methods=["GET"])
@tenant_required
def socio_list():
    args = request.args
    site_err = _enforce_site(args.get("center"))
    if site_err:
        return site_err
    try:
        return jsonify(_soc_list_db(args)), 200
    except Exception as exc:
        logger.warning("socio_list: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
    items = [r for r in _SOC_DB.values() if r.get("tenant_id") == _tenant()]
    if args.get("child_id"):
        items = [r for r in items if str(r.get("child_id")) == args["child_id"]]
    if args.get("center"):
        items = [r for r in items if r.get("center") == args["center"]]
    items.sort(key=lambda r: r.get("assessed_at", ""), reverse=True)
    page = max(int(args.get("page", 1)), 1)
    per_page = min(max(int(args.get("per_page", 20)), 1), 100)
    return jsonify(_paginate(items, page, per_page)), 200


@cmci_bp.route("/socioeconomic/<int:rec_id>", methods=["GET"])
@tenant_required
def socio_detail(rec_id):
    try:
        record = _soc_detail_db(rec_id)
    except Exception as exc:
        logger.warning("socio_detail: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
        err = _check_idor(_SOC_DB.get(rec_id))
        if err:
            return err
        return jsonify(_SOC_DB[rec_id]), 200
    if record is None:
        return jsonify({"error": "Registro no encontrado"}), 404
    return jsonify(record), 200


# --------------------------- PARAMS ---------------------------

@cmci_bp.route("/params", methods=["GET"])
@tenant_required
def params_get():
    scope = _scope()
    return jsonify({"scope": scope, "version": STORE.version,
                    "params": STORE.get(scope)}), 200


@cmci_bp.route("/params", methods=["PUT"])
@tenant_required
def params_put():
    """Solo license_admin. Scope global inmutable (seed validada)."""
    denied = _require_admin()
    if denied:
        return denied
    body = request.get_json(force=True, silent=True) or {}
    scope = body.get("scope", "global")
    patch = body.get("patch", {})
    try:
        updated = STORE.set(scope, patch,
                            actor=getattr(getattr(g, "current_user", None), "id", None))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 422
    return jsonify({"scope": scope, "version": STORE.version, "params": updated}), 200


# --------------------------- IMPORT EXCELS ---------------------------

@cmci_bp.route("/import-excels", methods=["POST"])
@tenant_required
def import_excels():
    """Lee BASE DE DATOS de ambos Excels (openpyxl, valores cacheados) como
    snapshots históricos. Recomputa Y/Z sobre lo importado."""
    denied = _deny_unless_can_create()
    if denied:
        return denied
    try:
        import openpyxl
    except ImportError:
        return jsonify({"error": "openpyxl no instalado"}), 500
    import os as _os
    body = request.get_json(force=True, silent=True) or {}
    seed_dir = _os.path.abspath(_os.path.join(_os.path.dirname(__file__), "..", "seeds", "cmci"))
    vuln_path = body.get("vuln_path") or _os.path.join(
        seed_dir, "Matriz_Vulnerabilidad_CMCI_VALIDADA_VERSION_1_OK_pruebas.xlsm")
    socio_path = body.get("socio_path") or _os.path.join(
        seed_dir, "Ficha_Socioeconomica_CMCI_v2_PROTOTIPO.xlsx")

    vuln_rows, soc_rows = [], []
    wb = openpyxl.load_workbook(vuln_path, data_only=True, read_only=True)
    ws = wb["BASE DE DATOS"]
    for row in ws.iter_rows(min_row=5, max_row=152, values_only=True):
        code = row[0]
        if not code:
            continue
        total = row[15] or 0  # P
        alert = (row[19] == "Sí")  # T
        prio = row[18] or ("PRIORIDAD 1" if alert else "PRIORIDAD 3")  # S
        result = {"total": float(total), "pct": (row[16] or 0),
                  "level": row[17], "priority": prio,
                  "protection_flag": alert, "protection_alert": "SÍ" if alert else "NO",
                  "semaphore": row[21], "childcare_need": row[20],
                  "subtotals": {"D1": row[7], "D2": row[8], "D3": row[9], "D4": row[10],
                                "D5": row[11], "D6": row[12], "D7": row[13], "D8": row[14]},
                  "order_y": VE.order_y(prio, alert, float(total)),
                  "scores": {}, "contrib": {}, "alerts": {}}
        assessed = row[6]
        vuln_rows.append({"body": {"code": code, "child_name": row[1],
                                   "center": row[4],
                                   "status": row[22] or "Validada",
                                   "assessed_at": assessed,
                                   "observations": row[23]},
                          "answers": {}, "result": result})
    wb.close()

    wb2 = openpyxl.load_workbook(socio_path, data_only=True, read_only=True)
    ws2 = wb2["BASE_DATOS"]
    for row in ws2.iter_rows(min_row=4, max_row=203, values_only=True):
        if not row[0]:
            continue
        assessed = row[1]
        soc_rows.append({"body": {"code": row[0], "center": row[2],
                                  "child_name": row[3], "assessed_at": assessed},
                         "data": {"integrantes": row[5]},
                         "result": {"ingreso_total": row[6], "gasto_total": row[7],
                                    "per_capita": row[8], "total": row[13],
                                    "classification": row[14]}})
    wb2.close()

    imported = {"vulnerability": 0, "socioeconomic": 0}
    try:
        for item in vuln_rows:
            _vuln_insert_db(item["body"], item["answers"], item["result"],
                            STORE.version, "excel-import")
            imported["vulnerability"] += 1
        for item in soc_rows:
            _soc_insert_db(item["body"], item["data"], item["result"],
                           STORE.version, "excel-import")
            imported["socioeconomic"] += 1
    except Exception as exc:
        logger.warning("import_excels: DB no disponible (%s); fallback memoria", exc)
        _rollback_quiet()
        for item in vuln_rows:
            body = item["body"]
            if hasattr(body.get("assessed_at"), "isoformat"):
                body = dict(body, assessed_at=body["assessed_at"].isoformat())
            elif body.get("assessed_at"):
                body = dict(body, assessed_at=str(body["assessed_at"]))
            else:
                body = dict(body, assessed_at=_now_iso())
            _SEQ["vuln"] += 1
            rec_id = _SEQ["vuln"]
            _VULN_DB[rec_id] = {"id": rec_id, "tenant_id": _tenant(),
                                "code": body["code"],
                                "child_name": body["child_name"],
                                "center": body["center"],
                                "status": body["status"],
                                "assessed_at": body["assessed_at"],
                                "observations": body["observations"],
                                "answers": {}, "result": item["result"],
                                "params_version": STORE.version,
                                "source": "excel-import", "created_by": None}
            imported["vulnerability"] += 1
        for item in soc_rows:
            body = item["body"]
            if hasattr(body.get("assessed_at"), "isoformat"):
                body = dict(body, assessed_at=body["assessed_at"].isoformat())
            elif body.get("assessed_at"):
                body = dict(body, assessed_at=str(body["assessed_at"]))
            else:
                body = dict(body, assessed_at=_now_iso())
            _SEQ["soc"] += 1
            rec_id = _SEQ["soc"]
            _SOC_DB[rec_id] = {"id": rec_id, "tenant_id": _tenant(),
                               "code": body["code"],
                               "center": body["center"],
                               "child_name": body["child_name"],
                               "assessed_at": body["assessed_at"],
                               "data": item["data"], "result": item["result"],
                               "params_version": STORE.version,
                               "source": "excel-import", "created_by": None}
            imported["socioeconomic"] += 1
    return jsonify({"imported": imported, "params_version": STORE.version}), 201


# --------------------------- UPDATE FICHAS (educador/coordinador/propietario) ---

@cmci_bp.route("/vulnerability/<int:rec_id>", methods=["PUT"])
@tenant_required
def vuln_update(rec_id):
    """Edita una ficha: recalcula puntaje y propaga a Child/Family.

    Roles: educadora/auxiliar (crean) + coordinadora + admin override.
    Recalcula con el motor determinista y actualiza la fila + datos del niño.
    """
    denied = _deny_unless_can_edit()
    if denied:
        return denied
    from models import db
    from models.cmci import VulnerabilityAssessment

    row = (VulnerabilityAssessment.query
           .filter(VulnerabilityAssessment.id == rec_id,
                   VulnerabilityAssessment.center_id == _tenant()).first())
    if row is None:
        return jsonify({"error": "Registro no encontrado"}), 404
    body = request.get_json(force=True, silent=True) or {}
    answers = _normalize_phones(body.get("answers", {}))
    # si no mandan answers, conserva las guardadas y solo actualiza meta
    if not answers:
        stored = dict(getattr(row, "answers", None) or {})
        stored.pop(_META, None)
        answers = stored
    try:
        params = STORE.get(body.get("scope", "global"))
        result = VE.compute(answers, params)
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    day, raw = _parse_day(body.get("assessed_at") or _now_iso())
    stored_answers = dict(answers)
    stored_answers[_META] = {
        "center": body.get("center"), "child_name": body.get("child_name"),
        "observations": body.get("observations"),
        "params_version": STORE.version, "source": "api-update",
        "assessed_at": raw, "result": result,
    }
    if body.get("child_id") is not None:
        try:
            row.child_id = int(body.get("child_id"))
        except (TypeError, ValueError):
            pass
    if body.get("code"):
        row.code = body.get("code")
    if day is not None:
        row.assessed_at = day
    bday_in = answers.get("birth_date")
    if bday_in:
        bday, _ = _parse_day(bday_in)
        if bday is not None:
            row.birth_date = bday
    row.age_months = result.get("age_months")
    row.in_range = (result.get("age_range") != "FUERA DE RANGO")
    row.answers = stored_answers
    row.scores = result.get("scores") or {}
    row.subtotals = result.get("subtotals") or {}
    row.total = result.get("total") or 0.0
    row.level = result.get("level")
    row.semaphore = result.get("semaphore")
    row.protection_alert = bool(result.get("protection_flag"))
    row.priority = result.get("priority")
    row.childcare_need = result.get("childcare_need")
    row.alerts = result.get("alerts") or {}
    if body.get("status"):
        row.status = body.get("status")
    db.session.add(row)
    db.session.commit()
    _sync_child_family_from_ficha({"child_id": row.child_id}, answers)
    return jsonify(_vuln_to_record(row)), 200


@cmci_bp.route("/socioeconomic/<int:rec_id>", methods=["PUT"])
@tenant_required
def socio_update(rec_id):
    """Edita ficha socioeconómica: recalcula SUMPRODUCT y propaga a Family."""
    denied = _deny_unless_can_edit()
    if denied:
        return denied
    from models import db
    from models.cmci import SocioeconomicAssessment

    row = (SocioeconomicAssessment.query
           .filter(SocioeconomicAssessment.id == rec_id,
                   SocioeconomicAssessment.center_id == _tenant()).first())
    if row is None:
        return jsonify({"error": "Registro no encontrado"}), 404
    body = request.get_json(force=True, silent=True) or {}
    data = _normalize_phones(body.get("data", {}))
    if not data:
        incomes = dict(getattr(row, "incomes", None) or {})
        meta = incomes.pop(_META, {})
        data = meta.get("data", {}) if isinstance(meta, dict) else {}
    try:
        result = SE.compute(data, STORE.get(body.get("scope", "global")))
    except (ValueError, KeyError) as exc:
        return jsonify({"error": str(exc)}), 422
    day, raw = _parse_day(body.get("assessed_at") or _now_iso())
    stored_in = dict(data.get("ingresos", {})) if isinstance(data.get("ingresos"), dict) else {}
    stored_in[_META] = {
        "center": body.get("center"), "child_name": body.get("child_name"),
        "code": body.get("code"), "params_version": STORE.version,
        "source": "api-update", "assessed_at": raw,
        "data": data, "result": result,
    }
    if body.get("child_id") is not None:
        try:
            row.child_id = int(body.get("child_id"))
        except (TypeError, ValueError):
            pass
    if body.get("code"):
        row.code = body.get("code")
    if day is not None:
        row.assessed_at = day
    row.incomes = stored_in
    if isinstance(data.get("egresos"), dict):
        row.expenses = dict(data.get("egresos"))
    row.per_capita = result.get("per_capita")
    row.expense_ratio = result.get("ratio_gasto")
    row.coverage_services = result.get("cobertura_servicios")
    row.dependency = result.get("personas_por_perceptor")
    row.subscores = result.get("subscores") or {}
    row.total = result.get("total") or 0.0
    row.classification = result.get("classification")
    db.session.add(row)
    db.session.commit()
    _sync_child_family_from_ficha({"child_id": row.child_id}, data)
    return jsonify(_soc_to_record(row)), 200


# --------------------------- BIBLIOTECA DOCUMENTAL CMCI ---------------------------
# 15 plantillas fijas §10.6. Subida: solo propietario (license_admin/
# super_admin) o coordinador principal (coordinator). Descarga/listado:
# cualquier rol autenticado del tenant (educadores, auxiliares, etc. solo
# ven botón Descargar). No se suben expedientes pesados, solo plantillas.

_BIB_ALLOWED_EXTS = {"pdf", "doc", "docx", "txt", "xls", "xlsx", "png", "jpg", "jpeg"}

# Mapeo slugs frontend (con guiones, params.ts) -> categorías backend (enum §10.6)
_BIB_SLUG_MAP = {
    "protocolo-ingreso": "protocolo_requisitos_ingreso",
    "ficha-postulacion": "ficha_postulacion",
    "ficha-cdp": "ficha_cdp",
    "ficha-socioeconomica": "ficha_socioeconomica",
    "ficha-vulnerabilidad": "ficha_vulnerabilidad",
    "informe-visita": "informe_tecnico_visita",
    "acta-compromiso": "acta_compromiso_corresponsabilidad",
    "consentimiento": "consentimiento_informado",
    "autorizacion-imagen": "autorizacion_imagen",
    "ficha-idii": "ficha_idii",
    "historia-clinica": "historia_clinica",
    "monitoreo-nutricional": "monitoreo_nutricional_curvas",
    "ficha-alimentacion": "ficha_diaria_alimentacion",
    "menu-semanal": "menu_semanal",
    "informe-mensual": "informe_mensual",
}


def _bib_normalize_category(raw):
    s = (raw or "").strip()
    if s in _BIB_SLUG_MAP:
        return _BIB_SLUG_MAP[s]
    s2 = s.replace("-", "_")
    if s2 in _BIB_SLUG_MAP.values():
        return s2
    return s


def _biblio_to_dict(t):
    return {
        "id": t.id,
        "title": t.title,
        "category": t.category,
        "file_url": t.file_url,
        "version": t.version,
    }


@cmci_bp.route("/biblioteca", methods=["GET"])
@tenant_required
def biblioteca_list():
    """Lista plantillas. ?category= para filtrar. Todos los roles pueden ver
    y descargar; la subida/borrado se restringen en sus endpoints."""
    from models.cmci import DocumentTemplate

    q = DocumentTemplate.query.order_by(DocumentTemplate.id.asc())
    cat = request.args.get("category")
    if cat:
        q = q.filter(DocumentTemplate.category == _bib_normalize_category(cat))
    items = q.all()
    return jsonify({"templates": [_biblio_to_dict(t) for t in items],
                    "total": len(items)}), 200


@cmci_bp.route("/biblioteca/upload", methods=["POST"])
@tenant_required
def biblioteca_upload():
    """Sube una plantilla. Solo propietario o coordinador principal."""
    denied = _deny_unless_can_upload_biblioteca()
    if denied:
        return denied
    from flask import current_app
    from werkzeug.utils import secure_filename
    import os as _os
    import uuid as _uuid
    from models import db
    from models.cmci import DocumentTemplate

    title = (request.form.get("title") or "").strip()
    category = _bib_normalize_category(request.form.get("category"))
    version = (request.form.get("version") or "v1.0").strip()
    if not title:
        return jsonify({"error": "El título es obligatorio"}), 400
    if category not in DocumentTemplate.CATEGORIES:
        return jsonify({"error": "Categoría inválida. Use una de las 15 oficiales",
                        "valid": list(DocumentTemplate.CATEGORIES)}), 422
    if "file" not in request.files:
        return jsonify({"error": "No se envió ningún archivo"}), 400
    fh = request.files["file"]
    if not fh or not fh.filename:
        return jsonify({"error": "Nombre de archivo vacío"}), 400
    ext = fh.filename.rsplit(".", 1)[-1].lower() if "." in fh.filename else ""
    if ext not in _BIB_ALLOWED_EXTS:
        return jsonify({"error": f"Tipo no permitido .{ext}. "
                                 f"Permitidos: {', '.join(sorted(_BIB_ALLOWED_EXTS))}"}), 400
    base = _os.path.join(current_app.root_path, "static", "uploads", "biblioteca")
    _os.makedirs(base, exist_ok=True)
    unique = f"{_uuid.uuid4().hex}_{secure_filename(fh.filename)}"
    full = _os.path.join(base, unique)
    fh.save(full)
    tpl = DocumentTemplate(title=title, category=category,
                           file_url=f"/static/uploads/biblioteca/{unique}",
                           version=version)
    db.session.add(tpl)
    db.session.commit()
    logger.info("Biblioteca plantilla subida: '%s' (%s) por user %s",
                title, category, getattr(getattr(g, "current_user", None), "id", None))
    return jsonify({"message": "Plantilla cargada exitosamente",
                    "template": _biblio_to_dict(tpl)}), 201


@cmci_bp.route("/biblioteca/<int:tpl_id>/download", methods=["GET"])
@tenant_required
def biblioteca_download(tpl_id):
    """Descarga una plantilla. Todos los roles autenticados (incl. roles
    inferiores: solo ven este botón)."""
    import os as _os
    from flask import current_app, send_file
    from models.cmci import DocumentTemplate

    tpl = DocumentTemplate.query.filter_by(id=tpl_id).first()
    if not tpl or not tpl.file_url:
        return jsonify({"error": "Plantilla no encontrada"}), 404
    rel = tpl.file_url.split("/static/", 1)[-1] if "/static/" in tpl.file_url else tpl.file_url.lstrip("/")
    full = _os.path.join(current_app.root_path, "static", rel)
    if not _os.path.exists(full):
        return jsonify({"error": "Archivo no encontrado en el servidor"}), 404
    return send_file(full, as_attachment=True,
                     download_name=f"{tpl.category}_{tpl.title}".replace(" ", "_") +
                     "." + full.rsplit(".", 1)[-1].lower())


@cmci_bp.route("/biblioteca/<int:tpl_id>", methods=["DELETE"])
@tenant_required
def biblioteca_delete(tpl_id):
    """Elimina plantilla. Solo quien puede subir (propietario/coordinador)."""
    denied = _deny_unless_can_upload_biblioteca()
    if denied:
        return denied
    import os as _os
    from flask import current_app
    from models import db
    from models.cmci import DocumentTemplate

    tpl = DocumentTemplate.query.filter_by(id=tpl_id).first()
    if not tpl:
        return jsonify({"error": "Plantilla no encontrada"}), 404
    try:
        if tpl.file_url and "/static/uploads/biblioteca/" in tpl.file_url:
            rel = tpl.file_url.split("/static/", 1)[-1]
            full = _os.path.join(current_app.root_path, "static", rel)
            if _os.path.exists(full):
                _os.remove(full)
    except Exception as exc:
        logger.warning("biblioteca_delete: no se pudo borrar archivo (%s)", exc)
    db.session.delete(tpl)
    db.session.commit()
    return jsonify({"message": "Plantilla eliminada"}), 200
