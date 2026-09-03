"""
Export Service
Handles report data conversion to different formats: PDF, CSV, Excel, Google Sheets
"""
import csv
import io
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# Optional: openpyxl for Excel export
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class ExportService:
    """Service for exporting report data to different formats"""
    
    @staticmethod
    def _resolve_row_value(row: dict, header: str):
        """Resolve a row value by trying multiple key variations for a header"""
        key = header.lower().replace(' ', '_').replace('(', '').replace(')', '')
        value = row.get(key) or row.get(header.lower()) or row.get(header)
        if value is None:
            for k, v in row.items():
                if k.lower().startswith(header.lower()[:3]):
                    value = v
                    break
        return value

    @staticmethod
    def _strip_empty_columns(headers: list, raw_rows: list) -> tuple:
        """Remove columns where ALL data values are empty or dash"""
        if not raw_rows:
            return headers, raw_rows
        
        # Build column data for each header
        cols_to_keep = []
        for i, h in enumerate(headers):
            has_data = False
            for row in raw_rows:
                val = row[i] if isinstance(row, list) else None
                if val and str(val).strip() not in ('', '-', 'None'):
                    has_data = True
                    break
            if has_data:
                cols_to_keep.append(i)
        
        if len(cols_to_keep) == len(headers):
            return headers, raw_rows  # Nothing to strip
        
        new_headers = [headers[i] for i in cols_to_keep]
        new_rows = []
        for row in raw_rows:
            new_rows.append([row[i] for i in cols_to_keep])
        return new_headers, new_rows

    @staticmethod
    def _get_col_widths(headers: list, total_width: float) -> list:
        """Calculate proportional column widths based on header type"""
        # Weight map: wider for text-heavy columns, narrower for short data
        weight_map = {
            'centro': 2.0,
            'niño': 2.2,
            'observaciones': 2.5,
            'notas': 2.2,
            'indicador': 2.0,
            'diagnóstico': 1.8,
            'diagnostico': 1.8,
            'área': 1.5,
            'fecha': 1.1,
            'estado': 1.0,
            'tipo': 1.0,
            'nivel': 1.0,
            'hora entrada': 1.0,
            'hora salida': 1.0,
            'métrica': 2.0,
            'valor': 2.0,
            'n.': 0.4,
            'cédula': 1.0,
            'edad': 0.8,
            'peso (kg)': 0.9,
            'talla (cm)': 0.9,
            'z-peso': 0.8,
            'z-talla': 0.8,
            'imc': 0.8
        }
        
        weights = []
        for h in headers:
            h_lower = h.lower()
            w = weight_map.get(h_lower, 1.2)  # Default weight
            weights.append(w)
        
        total_weight = sum(weights)
        return [(w / total_weight) * total_width for w in weights]

    @staticmethod
    def to_pdf(report_data: dict, center_name: str = "Centro") -> io.BytesIO:
        """
        Generate PDF from report data
        Args:
            report_data: { title, period, headers, rows, summary }
            center_name: Name of the center for the header
        Returns: BytesIO buffer containing the PDF
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), 
                                leftMargin=0.5*inch, rightMargin=0.5*inch,
                                topMargin=0.5*inch, bottomMargin=0.5*inch)
        elements = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=18,
            spaceAfter=6
        )
        
        cell_style = ParagraphStyle(
            'CellStyle',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            alignment=1,  # CENTER
        )
        
        cell_style_left = ParagraphStyle(
            'CellStyleLeft',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            alignment=0,  # LEFT (better for names)
        )
        
        header_cell_style = ParagraphStyle(
            'HeaderCellStyle',
            parent=styles['Normal'],
            fontSize=9,
            leading=11,
            alignment=1,  # CENTER
            textColor=colors.whitesmoke,
            fontName='Helvetica-Bold',
        )
        
        # Header
        elements.append(Paragraph(report_data.get('title', 'Reporte'), title_style))
        elements.append(Paragraph(f"Centro: {center_name}", styles['Normal']))
        elements.append(Paragraph(f"Período: {report_data.get('period', '')}", styles['Normal']))
        elements.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 12))
        
        # Summary section if exists
        summary = report_data.get('summary', {})
        if summary:
            elements.append(Paragraph("<b>Resumen:</b>", styles['Normal']))
            for key, value in summary.items():
                if isinstance(value, dict):
                    for k, v in value.items():
                        elements.append(Paragraph(f"  • {k}: {v}", styles['Normal']))
                else:
                    friendly_key = key.replace('_', ' ').title()
                    elements.append(Paragraph(f"  • {friendly_key}: {value}", styles['Normal']))
            elements.append(Spacer(1, 12))
        
        # Data table
        headers = report_data.get('headers', [])
        rows = report_data.get('rows', [])
        
        if headers and rows:
            # Step 1: Convert dict rows to flat lists
            flat_rows = []
            for row in rows:
                if isinstance(row, dict):
                    row_values = []
                    for h in headers:
                        value = ExportService._resolve_row_value(row, h)
                        row_values.append(str(value) if value is not None else '')
                    flat_rows.append(row_values)
                else:
                    flat_rows.append([str(v) for v in row])
            
            # Step 2: Strip entirely empty columns (e.g., Observaciones with no data)
            headers, flat_rows = ExportService._strip_empty_columns(headers, flat_rows)
            
            # Step 3: Identify text-heavy columns for left alignment
            text_heavy = {'niño', 'observaciones', 'notas', 'indicador', 'centro', 'área'}
            text_heavy_indices = set()
            for i, h in enumerate(headers):
                if h.lower() in text_heavy:
                    text_heavy_indices.add(i)
            
            # Step 4: Wrap all cells in Paragraph objects for text wrapping
            header_row = [Paragraph(h, header_cell_style) for h in headers]
            table_data = [header_row]
            
            for row_vals in flat_rows:
                wrapped = []
                for i, val in enumerate(row_vals):
                    style = cell_style_left if i in text_heavy_indices else cell_style
                    wrapped.append(Paragraph(str(val), style))
                table_data.append(wrapped)
            
            # Step 5: Calculate proportional column widths
            total_width = 10 * inch  # Landscape usable width
            col_widths = ExportService._get_col_widths(headers, total_width)
            
            t = Table(table_data, colWidths=col_widths, repeatRows=1)
            t.setStyle(TableStyle([
                # Header styling
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4A5568')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('TOPPADDING', (0, 0), (-1, 0), 8),
                
                # Data rows styling
                ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
                ('TOPPADDING', (0, 1), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                
                # Alternating row colors
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
                
                # Grid
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('BOX', (0, 0), (-1, -1), 1, colors.black),
            ]))
            elements.append(t)
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def to_csv(report_data: dict) -> io.BytesIO:
        """
        Generate CSV from report data
        Returns: BytesIO buffer containing the CSV
        """
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        
        headers = report_data.get('headers', [])
        rows = report_data.get('rows', [])
        
        # Write header
        if headers:
            writer.writerow(headers)
        
        # Write data rows
        for row in rows:
            if isinstance(row, dict):
                # Convert dict to list matching headers
                row_values = []
                for h in headers:
                    key = h.lower().replace(' ', '_').replace('(', '').replace(')', '')
                    value = row.get(key) or row.get(h.lower()) or row.get(h)
                    if value is None:
                        for k, v in row.items():
                            if k.lower().startswith(h.lower()[:3]):
                                value = v
                                break
                    row_values.append(value if value is not None else '')
                writer.writerow(row_values)
            else:
                writer.writerow(row)
        
        # Convert to bytes
        output = io.BytesIO()
        output.write(buffer.getvalue().encode('utf-8-sig'))  # BOM for Excel compatibility
        output.seek(0)
        return output
    
    @staticmethod
    def to_excel(report_data: dict, center_name: str = "Centro") -> io.BytesIO:
        """
        Generate Excel file from report data
        Returns: BytesIO buffer containing the Excel file
        """
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title and metadata
        ws['A1'] = report_data.get('title', 'Reporte')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"Centro: {center_name}"
        ws['A3'] = f"Período: {report_data.get('period', '')}"
        ws['A4'] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        
        start_row = 6
        
        # Headers
        headers = report_data.get('headers', [])
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Data rows
        rows = report_data.get('rows', [])
        for row_idx, row in enumerate(rows, start_row + 1):
            if isinstance(row, dict):
                for col, h in enumerate(headers, 1):
                    key = h.lower().replace(' ', '_').replace('(', '').replace(')', '')
                    value = row.get(key) or row.get(h.lower()) or row.get(h)
                    if value is None:
                        for k, v in row.items():
                            if k.lower().startswith(h.lower()[:3]):
                                value = v
                                break
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="center")
            else:
                for col, value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="center")
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column].width = adjusted_width
        
        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def to_google_sheets(report_data: dict, credentials, title: str = None) -> str:
        """
        Export report data to a new Google Sheet
        Args:
            report_data: { title, period, headers, rows }
            credentials: Google OAuth credentials object
            title: Optional title for the spreadsheet
        Returns: URL of the created spreadsheet
        """
        from googleapiclient.discovery import build
        
        if not credentials:
            raise ValueError("Google credentials required for Sheets export")
        
        # Build services
        sheets_service = build('sheets', 'v4', credentials=credentials)
        drive_service = build('drive', 'v3', credentials=credentials)
        
        # Create spreadsheet
        spreadsheet_title = title or f"{report_data.get('title', 'Reporte')} - {datetime.now().strftime('%Y-%m-%d')}"
        
        spreadsheet_body = {
            'properties': {'title': spreadsheet_title},
            'sheets': [{
                'properties': {
                    'title': 'Datos',
                    'gridProperties': {
                        'frozenRowCount': 1  # Freeze header row
                    }
                }
            }]
        }
        
        spreadsheet = sheets_service.spreadsheets().create(body=spreadsheet_body).execute()
        spreadsheet_id = spreadsheet.get('spreadsheetId')
        
        # Prepare data
        headers = report_data.get('headers', [])
        rows = report_data.get('rows', [])
        
        values = [headers]  # Start with headers
        
        for row in rows:
            if isinstance(row, dict):
                row_values = []
                for h in headers:
                    key = h.lower().replace(' ', '_').replace('(', '').replace(')', '')
                    value = row.get(key) or row.get(h.lower()) or row.get(h)
                    if value is None:
                        for k, v in row.items():
                            if k.lower().startswith(h.lower()[:3]):
                                value = v
                                break
                    row_values.append(value if value is not None else '')
                values.append(row_values)
            else:
                values.append(list(row))
        
        # Update spreadsheet with data
        body = {'values': values}
        sheets_service.spreadsheets().values().update(
            spreadsheetId=spreadsheet_id,
            range='Datos!A1',
            valueInputOption='RAW',
            body=body
        ).execute()
        
        # Format header row
        requests = [{
            'repeatCell': {
                'range': {
                    'sheetId': 0,
                    'startRowIndex': 0,
                    'endRowIndex': 1
                },
                'cell': {
                    'userEnteredFormat': {
                        'backgroundColor': {'red': 0.29, 'green': 0.34, 'blue': 0.41},
                        'textFormat': {
                            'foregroundColor': {'red': 1, 'green': 1, 'blue': 1},
                            'bold': True
                        },
                        'horizontalAlignment': 'CENTER'
                    }
                },
                'fields': 'userEnteredFormat(backgroundColor,textFormat,horizontalAlignment)'
            }
        }]
        
        sheets_service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={'requests': requests}
        ).execute()
        
        return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}"